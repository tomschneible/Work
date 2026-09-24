import cv2
import openpyxl
import pytest

from answer_extractor.detect import QuestionResult
from answer_extractor.export import LOW_CONFIDENCE_FILL, PATTERN_INFERRED_FILL, write_xlsx
from answer_extractor.pipeline import SheetResult, process_path, process_paths, untaken_sections
from answer_extractor.template import Template
from tests.synth import render_oval_sheet, render_sheet


def make_template() -> Template:
    data = {
        "page": {"width": 900, "height": 700},
        "sections": [
            {
                "name": "Answers",
                "columns": [
                    {"first_question": 1, "last_question": 4, "x_start": 150, "y_start": 100, "row_height": 80},
                ],
            }
        ],
        "bubble_spacing_x": 60,
        "bubble_radius": 18,
        "choices": {"even": ["A", "B", "C", "D"], "odd": ["F", "G", "H", "J"]},
        "thresholds": {"fill_ratio_min": 0.35, "relative_margin": 0.15},
    }
    return Template.from_dict(data)


def test_process_path_reads_image_files(tmp_path):
    template = make_template()
    image = render_sheet(template, {1: ["F"], 2: ["B"], 3: [], 4: ["A", "D"]})
    image_path = tmp_path / "sheet_001.png"
    cv2.imwrite(str(image_path), image)

    results = process_path(image_path, template)
    assert len(results) == 1
    result = results[0]
    assert result.label == "sheet_001"
    answers = {q.question: q.answer for q in result.questions}
    assert answers == {1: "F", 2: "B", 3: "", 4: "MULTIPLE"}
    assert result.has_review_items


def test_process_path_reads_directory_of_images(tmp_path):
    template = make_template()
    for i, marks in enumerate([{1: ["F"]}, {1: ["G"]}], start=1):
        image = render_sheet(template, marks)
        cv2.imwrite(str(tmp_path / f"sheet_{i}.png"), image)

    results = process_path(tmp_path, template)
    assert len(results) == 2
    assert {r.label for r in results} == {"sheet_1", "sheet_2"}


def test_process_paths_combines_multiple_inputs_and_dedupes_labels(tmp_path):
    template = make_template()

    dir_a = tmp_path / "batch_a"
    dir_a.mkdir()
    cv2.imwrite(str(dir_a / "sheet.png"), render_sheet(template, {1: ["F"]}))

    dir_b = tmp_path / "batch_b"
    dir_b.mkdir()
    cv2.imwrite(str(dir_b / "sheet.png"), render_sheet(template, {1: ["G"]}))

    results = process_paths([dir_a, dir_b], template)
    assert len(results) == 2
    labels = {r.label for r in results}
    assert labels == {"sheet", "sheet_1"}, "colliding labels from different inputs should be disambiguated"

    by_label = {r.label: r for r in results}
    answers_a = {q.question: q.answer for q in by_label["sheet"].questions}
    answers_b = {q.question: q.answer for q in by_label["sheet_1"].questions}
    assert answers_a[1] == "F"
    assert answers_b[1] == "G"


def test_write_xlsx_produces_file(tmp_path):
    template = make_template()
    image = render_sheet(template, {1: ["F"], 2: ["B"], 3: [], 4: ["A", "D"]})
    image_path = tmp_path / "sheet_001.png"
    cv2.imwrite(str(image_path), image)

    results = process_path(image_path, template)
    out_path = tmp_path / "results.xlsx"
    write_xlsx(results, out_path)

    assert out_path.exists()
    assert out_path.stat().st_size > 0


def _written_cells(tmp_path, questions, fallback_sections=()):
    result = SheetResult(
        label="sheet",
        source="test",
        used_contour_alignment=False,
        questions=questions,
        fallback_sections=list(fallback_sections),
    )
    path = tmp_path / "flagged.xlsx"
    write_xlsx([result], path)
    return openpyxl.load_workbook(path).worksheets[0]


def test_an_answer_read_without_confidence_is_colored_and_explained(tmp_path):
    """The only thing that flagged a real report was one answer read without
    confidence -- the spreadsheet has to show it as clearly as any other
    reason, not only in gray italics."""
    ws = _written_cells(
        tmp_path,
        [
            QuestionResult("Science", 1, "C", ["C"], {}, low_confidence=False),
            QuestionResult("Science", 2, "C", ["C"], {}, low_confidence=True),
        ],
    )

    confident, unsure = ws["B2"], ws["B3"]
    assert confident.fill.fill_type is None and confident.comment is None
    assert unsure.value == "C"
    assert unsure.fill.start_color.rgb == LOW_CONFIDENCE_FILL.start_color.rgb
    assert "not with confidence" in unsure.comment.text
    assert unsure.font.italic


def test_a_more_specific_flag_keeps_its_own_color(tmp_path):
    ws = _written_cells(
        tmp_path, [QuestionResult("Science", 1, "C", ["C"], {}, low_confidence=True, pattern_inferred=True)]
    )

    assert ws["B2"].fill.start_color.rgb == PATTERN_INFERRED_FILL.start_color.rgb


def test_a_section_whose_bubbles_werent_located_has_its_heading_colored_and_explained(tmp_path):
    ws = _written_cells(
        tmp_path,
        [
            QuestionResult("English", 1, "A", ["A"], {}, low_confidence=False),
            QuestionResult("Science", 1, "C", ["C"], {}, low_confidence=False),
        ],
        fallback_sections=["Science"],
    )

    english, science = ws["B1"], ws["C1"]
    assert english.fill.fill_type is None and english.comment is None
    assert science.fill.start_color.rgb == LOW_CONFIDENCE_FILL.start_color.rgb
    assert "couldn't be located" in science.comment.text


def make_two_section_template(science_optional=True) -> Template:
    def column(y):
        return [{"first_question": 1, "last_question": 4, "x_start": 150, "y_start": y, "row_height": 60}]

    return Template.from_dict(
        {
            "page": {"width": 900, "height": 900},
            "sections": [
                {"name": "English", "columns": column(100)},
                {"name": "Science", "optional": science_optional, "columns": column(500)},
            ],
            "bubble_spacing_x": 60,
            "bubble_radius": 18,
            "choices": {"even": ["A", "B", "C", "D"], "odd": ["F", "G", "H", "J"]},
        }
    )


def _question(section, question, answer="", **flags):
    return QuestionResult(section, question, answer, [answer] if answer else [], {}, low_confidence=False, **flags)


def _sheet(english_answers, science):
    return [_question("English", q, a) for q, a in enumerate(english_answers, start=1)] + science


def test_an_optional_section_left_entirely_blank_is_untaken_and_does_not_flag(tmp_path):
    template = make_two_section_template()
    questions = _sheet("FBGA", [_question("Science", q) for q in range(1, 5)])

    untaken = untaken_sections(template, questions, fallback_sections=[])
    result = SheetResult("x", "x", False, questions, untaken_sections=untaken)

    assert untaken == ["Science"]
    assert not result.has_review_items


def test_an_untaken_section_does_not_hide_a_blank_elsewhere():
    template = make_two_section_template()
    questions = _sheet("FB A", [_question("Science", q) for q in range(1, 5)])
    questions[2] = _question("English", 3)  # a real blank in a required section

    result = SheetResult("x", "x", False, questions, untaken_sections=untaken_sections(template, questions, []))

    assert result.has_review_items


@pytest.mark.parametrize(
    "science",
    [
        [_question("Science", 1, "F")] + [_question("Science", q) for q in range(2, 5)],  # one answer: taken
        [QuestionResult("Science", 1, "", [], {}, low_confidence=True)]
        + [_question("Science", q) for q in range(2, 5)],
        [_question("Science", 1, unreadable=True)] + [_question("Science", q) for q in range(2, 5)],
    ],
    ids=["one-answer", "low-confidence", "unreadable"],
)
def test_an_optional_section_with_anything_in_it_is_not_untaken(science):
    template = make_two_section_template()
    assert untaken_sections(template, _sheet("FBGA", science), []) == []


def test_a_blank_optional_section_whose_grid_was_not_found_is_not_untaken():
    template = make_two_section_template()
    questions = _sheet("FBGA", [_question("Science", q) for q in range(1, 5)])
    assert untaken_sections(template, questions, fallback_sections=["Science"]) == []


def test_a_required_section_left_blank_is_never_untaken():
    template = make_two_section_template(science_optional=False)
    questions = _sheet("FBGA", [_question("Science", q) for q in range(1, 5)])
    assert untaken_sections(template, questions, []) == []


@pytest.mark.parametrize(
    "name, optional",
    [("act_answer_sheet", True), ("act_j_form_answer_sheet", True), ("legacy_act_answer_sheet", False)],
)
def test_science_is_optional_only_on_the_enhanced_templates(name, optional):
    template = Template.from_yaml(f"templates/{name}.yaml")
    assert {s.name: s.optional for s in template.sections} == {
        "English": False, "Mathematics": False, "Reading": False, "Science": optional
    }


def test_a_real_enhanced_sheet_with_science_left_blank_is_not_flagged(tmp_path):
    template = Template.from_yaml("templates/act_answer_sheet.yaml")
    answers = {k: bubbles[k[1] % len(bubbles)].choice for k, bubbles in template.bubbles().items() if k[0] != "Science"}
    cv2.imwrite(str(tmp_path / "sheet.png"), render_oval_sheet(template, answers))

    (result,) = process_path(tmp_path / "sheet.png", template)

    assert result.untaken_sections == ["Science"]
    assert not result.has_review_items
    assert {(q.section, q.question): q.answer for q in result.questions if q.section != "Science"} == answers
